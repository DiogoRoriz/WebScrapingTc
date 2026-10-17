"""Leitura e gravação das planilhas Excel que simulam o banco de dados.

As abas do modelo estrela (produto, loja, oferta) são gravadas no padrão
visual do dicionário de dados:

    linha 1   faixa colorida com o nome da tabela
    linha 2   tipo e chave (pk/fk) de cada coluna
    linha 3   cabeçalho com o nome das colunas
    linha 4+  registros, com linhas alternadas

Cada uma também vira uma Tabela do Excel (tb_produto, tb_loja, tb_oferta),
que o Power BI reconhece diretamente. O arquivo ganha ainda a aba
"dicionario", com a descrição de todas as colunas.

A gravação é feita primeiro num arquivo temporário e depois o arquivo final
é substituído de uma vez, para que uma falha no meio nunca deixe a planilha
corrompida. Se o arquivo estiver aberto no Excel (ou bloqueado pelo
OneDrive), o programa tenta novamente por alguns minutos.
"""

import logging
import os
import time
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import esquema

log = logging.getLogger(__name__)

TENTATIVAS = 10
ESPERA = 30  # segundos entre tentativas quando o arquivo está bloqueado
ABA_DICIONARIO = "dicionario"
LINHA_CABECALHO = 3  # linha do cabeçalho nas abas formatadas

# Paleta do dicionário de dados.
COR_CABECALHO = "3C3489"
COR_PK = "D85A30"
COR_FK = "0F6E56"
COR_LINHA_ALTERNADA = "F1EFE8"
COR_BORDA = "D3D1C7"
COR_TEXTO = "2C2C2A"
COR_TEXTO_SECUNDARIO = "5F5E5A"
COR_LINK = "185FA5"

_BORDA_LINHA = Border(bottom=Side(style="thin", color=COR_BORDA))
_FONTE_CODIGO = "Consolas"


class ArquivoBloqueado(Exception):
    """O arquivo continuou aberto em outro programa após as tentativas."""


# ----------------------------------------------------------------------
# Leitura
# ----------------------------------------------------------------------
def _linha_do_cabecalho(arquivo, aba, primeira_coluna):
    """Índice (base 0) da linha onde está o cabeçalho da aba."""
    amostra = pd.read_excel(arquivo, sheet_name=aba, header=None, nrows=10)
    if amostra.empty:
        return 0
    for indice, valor in enumerate(amostra.iloc[:, 0]):
        if str(valor).strip() == primeira_coluna:
            return indice
    return 0


def ler_aba(arquivo, aba, colunas):
    """Lê uma aba da planilha. Se o arquivo ou a aba não existirem, devolve
    um DataFrame vazio com as colunas informadas."""
    if not arquivo.exists():
        return pd.DataFrame(columns=colunas)
    try:
        cabecalho = _linha_do_cabecalho(arquivo, aba, colunas[0])
        tabela = pd.read_excel(arquivo, sheet_name=aba, header=cabecalho)
    except ValueError:  # aba inexistente
        return pd.DataFrame(columns=colunas)
    tabela = tabela.dropna(how="all")
    for coluna in colunas:
        if coluna not in tabela.columns:
            tabela[coluna] = pd.NA
    return tabela[colunas].reset_index(drop=True)


def _ler_todas(arquivo):
    """Lê todas as abas de dados de um arquivo existente."""
    if not arquivo.exists():
        return {}
    abas = {}
    with pd.ExcelFile(arquivo) as livro:  # fecha o arquivo ao terminar
        nomes = livro.sheet_names
    for nome in nomes:
        if nome == ABA_DICIONARIO:
            continue  # é recriada a cada gravação
        if nome in esquema.TABELAS:
            abas[nome] = ler_aba(arquivo, nome, esquema.colunas(nome))
        else:
            abas[nome] = pd.read_excel(arquivo, sheet_name=nome)
    return abas


# ----------------------------------------------------------------------
# Gravação
# ----------------------------------------------------------------------
def salvar_abas(arquivo, abas):
    """Grava várias abas ({nome: DataFrame}) em um único arquivo .xlsx,
    preservando as abas existentes que não foram informadas."""
    arquivo.parent.mkdir(parents=True, exist_ok=True)
    existentes = _ler_todas(arquivo)
    existentes.update(abas)

    livro = Workbook()
    livro.remove(livro.active)
    do_modelo = [n for n in esquema.TABELAS if n in existentes]
    if do_modelo:
        _escrever_dicionario(livro.create_sheet(ABA_DICIONARIO))
    for nome in do_modelo:
        _escrever_tabela_modelo(livro.create_sheet(nome), nome,
                                existentes[nome])
    for nome, tabela in existentes.items():
        if nome not in esquema.TABELAS:
            _escrever_tabela_simples(livro.create_sheet(nome), tabela)

    temporario = arquivo.with_name(f"~tmp_{arquivo.name}")
    try:
        livro.save(temporario)
    except Exception:
        temporario.unlink(missing_ok=True)
        raise

    for tentativa in range(1, TENTATIVAS + 1):
        try:
            os.replace(temporario, arquivo)
            return
        except PermissionError:
            log.warning("%s está aberto em outro programa (tentativa %s/%s)."
                        " Feche o arquivo no Excel.", arquivo.name,
                        tentativa, TENTATIVAS)
            time.sleep(ESPERA)
    raise ArquivoBloqueado(
        f"Não foi possível gravar {arquivo}: o arquivo continua aberto. "
        f"Os dados desta execução ficaram em {temporario}.")


def anexar_linhas(arquivo, aba, linhas, colunas):
    """Acrescenta linhas ao final de uma aba (criando-a se preciso)."""
    atual = ler_aba(arquivo, aba, colunas)
    novas = pd.DataFrame(linhas, columns=colunas)
    if atual.empty:
        combinado = novas
    else:
        combinado = pd.concat([atual, novas], ignore_index=True)
    salvar_abas(arquivo, {aba: combinado})


# ----------------------------------------------------------------------
# Formatação
# ----------------------------------------------------------------------
def _preencher(cor):
    return PatternFill("solid", start_color=cor, end_color=cor)


def _valor(valor):
    """Converte valores do pandas/numpy para tipos aceitos pelo Excel."""
    if valor is None:
        return None
    try:
        if pd.isna(valor):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(valor, pd.Timestamp):
        return valor.to_pydatetime()
    if hasattr(valor, "item"):  # escalares do numpy
        return valor.item()
    return valor


def _faixa_titulo(planilha, linha, coluna_inicial, coluna_final, texto, cor):
    planilha.merge_cells(start_row=linha, start_column=coluna_inicial,
                         end_row=linha, end_column=coluna_final)
    celula = planilha.cell(row=linha, column=coluna_inicial, value=texto)
    celula.font = Font(bold=True, size=16, color=COR_TEXTO)
    celula.alignment = Alignment(vertical="center", indent=1)
    for coluna in range(coluna_inicial, coluna_final + 1):
        planilha.cell(row=linha, column=coluna).fill = _preencher(cor)
    planilha.row_dimensions[linha].height = 34


def _celula_cabecalho(celula, texto):
    celula.value = texto
    celula.fill = _preencher(COR_CABECALHO)
    celula.font = Font(bold=True, size=11, color="FFFFFF")
    celula.alignment = Alignment(vertical="center", indent=1)


def _celula_chave(celula, chave, texto=None):
    """Marcação de chave: pk em laranja, fk em verde."""
    celula.value = texto or chave
    celula.fill = _preencher(COR_PK if chave == "pk" else COR_FK)
    celula.font = Font(bold=True, color="FFFFFF")
    celula.alignment = Alignment(horizontal="center", vertical="center")


def _escrever_dicionario(planilha):
    """Aba com um quadro por tabela: Coluna, Tipo, Descrição e Chave."""
    planilha.sheet_view.showGridLines = False
    for letra, largura in zip("ABCDE", (2, 28, 14, 58, 12)):
        planilha.column_dimensions[letra].width = largura

    planilha["B1"] = "DICIONÁRIO DE DADOS"
    planilha["B1"].font = Font(bold=True, size=13, color=COR_TEXTO)
    planilha["B2"] = (
        "Modelo estrela do pipeline de preços · atualizado em "
        f"{datetime.now():%d/%m/%Y %H:%M}")
    planilha["B2"].font = Font(italic=True, size=10,
                               color=COR_TEXTO_SECUNDARIO)

    linha = 4
    for definicao in esquema.TABELAS.values():
        _faixa_titulo(planilha, linha, 2, 5, definicao["titulo"],
                      definicao["cor"])
        linha += 1
        for coluna, texto in zip(range(2, 6),
                                 ("Coluna", "Tipo", "Descrição", "Chave")):
            _celula_cabecalho(planilha.cell(row=linha, column=coluna), texto)
        planilha.row_dimensions[linha].height = 26
        linha += 1

        for posicao, item in enumerate(definicao["colunas"]):
            fundo = _preencher(
                COR_LINHA_ALTERNADA if posicao % 2 else "FFFFFF")
            valores = (item.nome, item.tipo, item.descricao)
            fontes = (
                Font(name=_FONTE_CODIGO, bold=True, color=COR_TEXTO),
                Font(name=_FONTE_CODIGO, color=COR_TEXTO_SECUNDARIO),
                Font(color=COR_TEXTO),
            )
            for coluna, valor, fonte in zip(range(2, 5), valores, fontes):
                celula = planilha.cell(row=linha, column=coluna, value=valor)
                celula.font = fonte
                celula.fill = fundo
                celula.border = _BORDA_LINHA
                celula.alignment = Alignment(vertical="center", indent=1)
            celula_chave = planilha.cell(row=linha, column=5)
            celula_chave.fill = fundo
            celula_chave.border = _BORDA_LINHA
            if item.chave:
                _celula_chave(celula_chave, item.chave)
            planilha.row_dimensions[linha].height = 24
            linha += 1
        linha += 2  # espaço entre os quadros


def _escrever_tabela_modelo(planilha, nome, tabela):
    """Aba de dados no padrão do dicionário, como Tabela do Excel."""
    definicao = esquema.TABELAS[nome]
    colunas = definicao["colunas"]
    total = len(colunas)
    planilha.sheet_view.showGridLines = False

    _faixa_titulo(planilha, 1, 1, total, definicao["titulo"],
                  definicao["cor"])

    # Linha 2: tipo e chave de cada coluna.
    for indice, item in enumerate(colunas, start=1):
        celula = planilha.cell(row=2, column=indice)
        if item.chave:
            _celula_chave(celula, item.chave, f"{item.chave} · {item.tipo}")
        else:
            celula.value = item.tipo
            celula.font = Font(name=_FONTE_CODIGO, size=9,
                               color=COR_TEXTO_SECUNDARIO)
            celula.fill = _preencher(COR_LINHA_ALTERNADA)
            celula.alignment = Alignment(horizontal="center")
    planilha.row_dimensions[2].height = 18

    # Linha 3: cabeçalho.
    for indice, item in enumerate(colunas, start=1):
        _celula_cabecalho(planilha.cell(row=LINHA_CABECALHO, column=indice),
                          item.nome)
    planilha.row_dimensions[LINHA_CABECALHO].height = 26

    # Registros.
    registros = tabela[[c.nome for c in colunas]].itertuples(index=False)
    for posicao, registro in enumerate(registros):
        linha = LINHA_CABECALHO + 1 + posicao
        fundo = _preencher(COR_LINHA_ALTERNADA if posicao % 2 else "FFFFFF")
        for indice, (item, valor) in enumerate(zip(colunas, registro),
                                               start=1):
            celula = planilha.cell(row=linha, column=indice,
                                   value=_valor(valor))
            celula.fill = fundo
            celula.border = _BORDA_LINHA
            _formatar_valor(celula, item)

    ultima_linha = max(LINHA_CABECALHO + len(tabela), LINHA_CABECALHO + 1)
    ultima_coluna = get_column_letter(total)
    tabela_excel = Table(
        displayName=f"tb_{nome}",
        ref=f"A{LINHA_CABECALHO}:{ultima_coluna}{ultima_linha}")
    tabela_excel.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight1", showRowStripes=False)
    planilha.add_table(tabela_excel)
    planilha.freeze_panes = f"A{LINHA_CABECALHO + 1}"

    for indice, item in enumerate(colunas, start=1):
        amostra = tabela[item.nome].head(300).map(
            lambda v: len(str(v)) if not pd.isna(v) else 0)
        largura = max([len(item.nome) + 4, len(item.tipo) + 8]
                      + amostra.tolist())
        if item.tipo == "TIMESTAMP":
            largura = 20
        planilha.column_dimensions[get_column_letter(indice)].width = min(
            largura + 2, 60)


def _formatar_valor(celula, item):
    """Formato de exibição conforme o tipo da coluna."""
    if item.tipo == "INT64":
        celula.number_format = "0"
        celula.alignment = Alignment(horizontal="center")
        celula.font = Font(name=_FONTE_CODIGO, color=COR_TEXTO)
    elif item.tipo == "FLOAT64":
        celula.number_format = '"R$" #,##0.00'
        celula.font = Font(color=COR_TEXTO)
    elif item.tipo == "BOOL":
        celula.alignment = Alignment(horizontal="center")
        celula.font = Font(color=COR_TEXTO)
    elif item.tipo == "TIMESTAMP":
        celula.number_format = "dd/mm/yyyy hh:mm:ss"
        celula.alignment = Alignment(horizontal="center")
        celula.font = Font(color=COR_TEXTO)
    elif item.nome.startswith("url") and celula.value:
        celula.hyperlink = celula.value
        celula.font = Font(color=COR_LINK, underline="single")
    else:
        celula.font = Font(color=COR_TEXTO)


def _escrever_tabela_simples(planilha, tabela):
    """Abas auxiliares (coleta bruta, log): cabeçalho, filtro e larguras."""
    planilha.append([str(c) for c in tabela.columns])
    for registro in tabela.itertuples(index=False):
        planilha.append([_valor(v) for v in registro])
    for celula in planilha[1]:
        _celula_cabecalho(celula, celula.value)
    planilha.freeze_panes = "A2"
    if len(tabela.columns):
        planilha.auto_filter.ref = planilha.dimensions
    for indice, coluna in enumerate(tabela.columns, start=1):
        if pd.api.types.is_datetime64_any_dtype(tabela[coluna]):
            for (celula,) in planilha.iter_rows(min_row=2, min_col=indice,
                                                max_col=indice):
                celula.number_format = "dd/mm/yyyy hh:mm:ss"
        amostra = tabela[coluna].head(200).map(lambda v: len(str(v)))
        largura = max([len(str(coluna))] + amostra.tolist())
        planilha.column_dimensions[get_column_letter(indice)].width = min(
            max(largura + 2, 10), 60)
