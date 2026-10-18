"""Testes rápidos da correspondência de anúncios e do tratamento.

Execute com:  python testes.py
"""

import unittest
from datetime import datetime, timedelta

import pandas as pd

import config
from correspondencia import anuncio_corresponde, escolher_melhor_oferta
from lojas.base import Anuncio, converter_preco
from tratamento import (calcular_indicadores, identificar_marca,
                        padronizar_loja)


class TesteCorrespondencia(unittest.TestCase):

    def casa(self, titulo, obrig, proib=""):
        return anuncio_corresponde(titulo, [t for t in obrig.split(";") if t],
                                   [t for t in proib.split(";") if t])

    def test_modelo_exato_nao_casa_com_variantes(self):
        self.assertTrue(self.casa(
            "Processador AMD Ryzen 5 5600, 3.5GHz (4.4GHz Turbo), AM4",
            "ryzen 5 5600"))
        self.assertFalse(self.casa("Processador AMD Ryzen 5 5600X",
                                   "ryzen 5 5600"))
        self.assertFalse(self.casa("Processador AMD Ryzen 5 5600GT",
                                   "ryzen 5 5600"))
        self.assertFalse(self.casa("Processador AMD Ryzen 5 7600X3D, 6-Core",
                                   "ryzen 5 7600"))
        self.assertTrue(self.casa("Processador AMD Ryzen 7 7800X3D",
                                  "ryzen 7 7800x3d"))
        self.assertTrue(self.casa("Placa de Vídeo ASUS Dual RTX3060 12GB",
                                  "rtx 3060;12gb"))
        self.assertFalse(self.casa("Ryzen 55600", "ryzen 5 5600"))

    def test_hifen_e_maiusculas(self):
        self.assertTrue(self.casa("Processador Intel Core I5 - 12400F",
                                  "i5 12400f"))
        self.assertTrue(self.casa("Placa-Mãe ASUS TUF Gaming B550M-PLUS",
                                  "tuf gaming;b550m plus"))

    def test_unidade_com_ou_sem_espaco_e_alternativas(self):
        titulo = "Memória Kingston Fury Beast, 16 GB, 3200MHz, DDR4"
        self.assertTrue(self.casa(titulo, "fury beast;16gb;ddr4;3200mhz|3200"))
        self.assertTrue(self.casa("Memoria 16Gb Ddr4 3200 Fury Beast",
                                  "fury beast;16gb;ddr4;3200mhz|3200"))

    def test_kits_e_proibidas(self):
        self.assertFalse(self.casa(
            "Memória Fury Beast 16GB (2x8GB) DDR4 3200MHz",
            "fury beast;8gb;ddr4", "16gb;2 x 8gb"))
        self.assertFalse(self.casa("Placa de Vídeo RTX 5060 Ti 8GB",
                                   "rtx 5060", "5060 ti"))
        self.assertFalse(self.casa(
            "Placa de Vídeo Gigabyte RTX 5060 + Fonte MSI MAG A650BN",
            "rtx 5060", "fonte"))

    def test_escolhe_mais_barato_disponivel(self):
        produto = {"palavras_obrigatorias": "rtx 5060",
                   "palavras_proibidas": "5060 ti", "preco_minimo": 1000}
        anuncios = [
            Anuncio("Placa RTX 5060 MSI", 2300.0, "a"),
            Anuncio("Placa RTX 5060 Galax", 2100.0, "b", disponivel=False),
            Anuncio("Placa RTX 5060 Ti", 2000.0, "c"),
            Anuncio("Suporte RTX 5060", 50.0, "d"),
            Anuncio("Notebook Gamer RTX 5060", 1900.0, "e"),
            Anuncio("Placa RTX 5060 ASUS", 2250.0, "f"),
        ]
        melhor = escolher_melhor_oferta(anuncios, produto,
                                        config.PALAVRAS_PROIBIDAS_GLOBAIS)
        self.assertEqual(melhor.url, "f")

    def test_computador_montado_nao_vira_processador(self):
        # Caso real (Amazon, 15/10/2026): PC montado aceito como processador.
        produto = {"palavras_obrigatorias": "i5 14400f",
                   "palavras_proibidas": "", "preco_minimo": 500}
        anuncios = [
            Anuncio("PC Home Core i5-14400F H610 DDR4 GT 730 4GB Windows 11",
                    3752.50, "pc"),
            Anuncio("Kit Intel B760 Gamer DDR4 I5 14400F 14ª Geração",
                    2641.00, "kit"),
            Anuncio("Processador Intel Core i5-14400F LGA 1700", 1000.0,
                    "cpu"),
        ]
        melhor = escolher_melhor_oferta(
            anuncios[:2], produto, config.PALAVRAS_PROIBIDAS_GLOBAIS)
        self.assertIsNone(melhor)
        melhor = escolher_melhor_oferta(
            anuncios, produto, config.PALAVRAS_PROIBIDAS_GLOBAIS)
        self.assertEqual(melhor.url, "cpu")

    def test_preco_minimo_vazio(self):
        produto = {"palavras_obrigatorias": "a400",
                   "preco_minimo": float("nan")}
        melhor = escolher_melhor_oferta([Anuncio("SSD A400", 150.0, "x")],
                                        produto, [])
        self.assertIsNotNone(melhor)

    def test_converter_preco(self):
        self.assertEqual(converter_preco("R$\xa01.404,90"), 1404.90)
        self.assertEqual(converter_preco("R$ 939,99"), 939.99)
        self.assertEqual(converter_preco("1.404"), 1404.0)
        self.assertEqual(converter_preco(559.0), 559.0)
        self.assertIsNone(converter_preco(""))


class TesteTratamento(unittest.TestCase):

    def test_padronizacao(self):
        self.assertEqual(padronizar_loja(" kabum "), "KaBuM!")
        self.assertEqual(padronizar_loja("KABUM!"), "KaBuM!")
        self.assertEqual(padronizar_loja("terabyteshop"), "Terabyte")
        self.assertEqual(padronizar_loja("loja nova"), "Loja Nova")
        self.assertEqual(identificar_marca(
            "Placa de Vídeo GeForce RTX 5060 8GB"), "NVIDIA")
        self.assertEqual(identificar_marca(
            "Placa-mãe ASUS TUF Gaming B550M-Plus"), "ASUS")
        self.assertEqual(identificar_marca(
            "Processador AMD Ryzen 5 5600"), "AMD")
        self.assertEqual(identificar_marca("SSD Kingston NV3 1TB"),
                         "Kingston")

    def test_media_sem_lookahead_e_outlier(self):
        t0 = datetime(2026, 10, 1, 3, 0)
        precos = [100, 110, 200, 90, 80]  # 200 é outlier
        oferta = pd.DataFrame({
            "id_oferta": range(1, 6),
            "id_produto": [1] * 5,
            "id_loja": [1] * 5,
            "preco": precos,
            "preco_medio_historico": None,
            "outlier": None,
            "situacao_compra": None,
            "url_oferta": "u",
            "data_coleta": [t0 + timedelta(days=i) for i in range(5)],
        })
        r = calcular_indicadores(oferta)
        self.assertTrue(pd.isna(r.loc[0, "preco_medio_historico"]))
        self.assertEqual(r.loc[0, "situacao_compra"],
                         config.SITUACAO_SEM_HISTORICO)
        self.assertEqual(r.loc[1, "preco_medio_historico"], 100.0)
        self.assertEqual(r.loc[1, "situacao_compra"], config.SITUACAO_MEDIA)
        self.assertEqual(r.loc[2, "preco_medio_historico"], 105.0)
        self.assertTrue(r.loc[2, "outlier"])
        self.assertEqual(r.loc[2, "situacao_compra"], config.SITUACAO_EVITAR)
        # O outlier (200) não entra na média seguinte: (100 + 110) / 2.
        self.assertEqual(r.loc[3, "preco_medio_historico"], 105.0)
        self.assertEqual(r.loc[3, "situacao_compra"], config.SITUACAO_BOA)
        # (100 + 110 + 90) / 3 = 100 -> 80 está 20% abaixo.
        self.assertEqual(r.loc[4, "preco_medio_historico"], 100.0)
        self.assertEqual(r.loc[4, "situacao_compra"], config.SITUACAO_BOA)
        self.assertFalse(r["outlier"].drop(index=2).any())

    def test_lojas_do_mesmo_dia_nao_se_influenciam(self):
        t0 = datetime(2026, 10, 1, 3, 0)
        oferta = pd.DataFrame({
            "id_oferta": [1, 2, 3], "id_produto": [7, 7, 7],
            "id_loja": [1, 2, 1], "preco": [100, 150, 120],
            "preco_medio_historico": None, "outlier": None,
            "situacao_compra": None, "url_oferta": "u",
            "data_coleta": [t0, t0 + timedelta(minutes=7),
                            t0 + timedelta(days=1)],
        })
        r = calcular_indicadores(oferta)
        self.assertEqual(list(r["situacao_compra"][:2]),
                         [config.SITUACAO_SEM_HISTORICO] * 2)
        self.assertEqual(r.loc[2, "preco_medio_historico"], 125.0)


class NavegadorFalso:
    """Simula o Chrome: falha `falhas` vezes como se a janela tivesse sido
    fechada e depois devolve uma página de busca da Amazon."""

    PAGINA = (
        '<div data-component-type="s-search-result" data-asin="B000000001">'
        '<h2 aria-label="Processador AMD Ryzen 5 5600 BOX"></h2>'
        '<span class="a-price"><span class="a-offscreen">R$ 800,00</span>'
        '</span></div>')

    def __init__(self, falhas):
        self.falhas, self.reinicios, self.chamadas = falhas, 0, 0

    def abrir(self, url, seletor):
        from selenium.common.exceptions import InvalidSessionIdException
        self.chamadas += 1
        if self.falhas:
            self.falhas -= 1
            raise InvalidSessionIdException("invalid session id")
        return self.PAGINA

    def reiniciar(self):
        self.reinicios += 1


class TesteChromeFechado(unittest.TestCase):

    PRODUTOS = [
        {"nome_produto": f"Processador AMD Ryzen 5 5600 ({i})",
         "categoria_produto": "Processador", "termo_busca": "ryzen 5 5600",
         "palavras_obrigatorias": "ryzen 5 5600", "palavras_proibidas": "",
         "preco_minimo": 300}
        for i in range(3)
    ]

    def setUp(self):
        import extracao
        self.extracao = extracao
        self.pausa_original = extracao.pausa_educada
        extracao.pausa_educada = lambda *a, **k: None

    def tearDown(self):
        self.extracao.pausa_educada = self.pausa_original

    def test_reabre_o_chrome_uma_vez_e_continua(self):
        nav = NavegadorFalso(falhas=1)
        resultado = self.extracao.ScriptExtracao(
            self.PRODUTOS, ["amazon"])._coletar_amazon(nav)
        self.assertEqual(nav.reinicios, 1)
        self.assertEqual(resultado.status, "OK")
        self.assertEqual(len(resultado.registros), 3)

    def test_desiste_se_o_chrome_cair_de_novo(self):
        nav = NavegadorFalso(falhas=99)
        resultado = self.extracao.ScriptExtracao(
            self.PRODUTOS, ["amazon"])._coletar_amazon(nav)
        self.assertEqual(nav.reinicios, 1)
        self.assertEqual(nav.chamadas, 2)  # não insiste produto a produto
        self.assertEqual(resultado.status, "ERRO")


class TesteIntegracaoPlanilhas(unittest.TestCase):
    """Carga + tratamento em planilhas temporárias, sem acessar as lojas."""

    def setUp(self):
        import tempfile
        from pathlib import Path
        self.pasta = tempfile.TemporaryDirectory()
        base = Path(self.pasta.name)
        self.originais = (config.ARQUIVO_BANCO, config.PASTA_COLETAS,
                          config.ARQUIVO_LOG_EXECUCOES)
        config.ARQUIVO_BANCO = base / "banco_dados.xlsx"
        config.PASTA_COLETAS = base / "coletas"
        config.ARQUIVO_LOG_EXECUCOES = base / "log_execucoes.xlsx"

    def tearDown(self):
        (config.ARQUIVO_BANCO, config.PASTA_COLETAS,
         config.ARQUIVO_LOG_EXECUCOES) = self.originais
        self.pasta.cleanup()

    @staticmethod
    def coleta(dia, precos):
        linhas = []
        for loja, produto, preco in precos:
            linhas.append({
                "data_coleta": datetime(2026, 10, dia, 3, 0, 0),
                "nome_loja": loja, "categoria_produto": "Processador",
                "nome_produto": produto, "titulo_anuncio": produto + " BOX",
                "preco": preco, "preco_prazo": preco * 1.1,
                "url_oferta": f"https://exemplo/{loja}/{dia}",
            })
        return pd.DataFrame(linhas)

    def test_pipeline_em_planilhas(self):
        from carga import executar_carga
        from extracao import ResultadoLoja, ScriptExtracao
        from tratamento import executar_tratamento

        r5 = "Processador AMD Ryzen 5 5600"
        i5 = "Processador Intel Core i5-12400F"
        dia1 = self.coleta(1, [("kabum", r5, 800), ("amazon", r5, 820),
                               ("kabum", i5, 900)])
        dia2 = self.coleta(2, [("kabum", r5, 700), ("pichau", r5, 1500),
                               ("kabum", i5, 910)])
        dia3 = self.coleta(3, [("kabum", r5, 812)])

        resumo = executar_carga(dia1)
        self.assertEqual(resumo["ofertas"], 3)
        self.assertEqual(executar_carga(dia1)["ofertas"], 0)  # idempotente
        executar_carga(dia2)
        executar_carga(dia3)
        executar_tratamento()
        executar_tratamento()  # reexecutar não muda o resultado

        import esquema
        import planilhas
        from openpyxl import load_workbook

        # Estrutura visual: dicionário + abas com faixa, tipos e cabeçalho.
        livro = load_workbook(config.ARQUIVO_BANCO)
        self.assertEqual(livro.sheetnames,
                         ["dicionario", "produto", "loja", "oferta"])
        aba = livro["oferta"]
        self.assertEqual(aba["A1"].value, "OFERTA")
        self.assertEqual(aba["A2"].value, "pk · INT64")
        self.assertEqual(aba["B2"].value, "fk · INT64")
        self.assertEqual(aba["A3"].value, "id_oferta")
        self.assertEqual(aba["A3"].fill.start_color.rgb[-6:],
                         planilhas.COR_CABECALHO)
        self.assertEqual(aba["A2"].fill.start_color.rgb[-6:],
                         planilhas.COR_PK)
        self.assertEqual(aba["B2"].fill.start_color.rgb[-6:],
                         planilhas.COR_FK)
        self.assertIn("tb_oferta", aba.tables)
        self.assertEqual(aba.tables["tb_oferta"].ref, "A3:I10")
        dicionario = [c.value for c in livro["dicionario"]["B"] if c.value]
        for tabela in esquema.TABELAS.values():
            self.assertIn(tabela["titulo"], dicionario)
            for coluna in tabela["colunas"]:
                self.assertIn(coluna.nome, dicionario)

        banco = {nome: planilhas.ler_aba(config.ARQUIVO_BANCO, nome,
                                         esquema.colunas(nome))
                 for nome in esquema.TABELAS}
        self.assertEqual(list(banco["produto"].columns),
                         ["id_produto", "nome_produto", "categoria_produto",
                          "marca"])
        self.assertEqual(list(banco["loja"].columns),
                         ["id_loja", "nome_loja", "url_base"])
        self.assertEqual(list(banco["oferta"].columns),
                         ["id_oferta", "id_produto", "id_loja", "preco",
                          "preco_medio_historico", "outlier",
                          "situacao_compra", "url_oferta", "data_coleta"])
        self.assertEqual(sorted(banco["loja"]["nome_loja"]),
                         ["Amazon", "KaBuM!", "Pichau"])
        self.assertEqual(sorted(banco["produto"]["marca"]), ["AMD", "Intel"])

        oferta = banco["oferta"].merge(banco["produto"], on="id_produto")
        r5_ofertas = oferta[oferta["nome_produto"] == r5].sort_values(
            "data_coleta")
        self.assertEqual(len(oferta), 7)
        # Dia 2: média 810 (800 e 820); 1500 é outlier e fica fora do
        # histórico do dia 3: (800 + 820 + 700) / 3 = 773,33.
        pichau = r5_ofertas[r5_ofertas["preco"] == 1500].iloc[0]
        self.assertTrue(pichau["outlier"])
        self.assertAlmostEqual(r5_ofertas.iloc[-1]["preco_medio_historico"],
                               773.33, places=2)
        self.assertEqual(r5_ofertas.iloc[-1]["situacao_compra"],
                         config.SITUACAO_MEDIA)

        # Planilha bruta: o que uma execução interrompida deixou gravado é
        # recuperado pela próxima carga, sem duplicar o que já foi carregado.
        from datetime import date
        from extracao import coleta_para_carga
        dia4 = self.coleta(4, [("amazon", i5, 950)])
        res = ResultadoLoja(loja="Amazon", registros=dia4.to_dict("records"))
        ScriptExtracao.salvar_tabela_unica([res], dia=date(2026, 10, 4))
        self.assertEqual(len(list(config.PASTA_COLETAS.glob("*.xlsx"))), 1)
        # Cópia em CSV para a ingestão da equipe, com as colunas do contrato.
        from extracao import COLUNAS_COLETA
        csv = pd.read_csv(config.PASTA_COLETAS / "coleta_2026-10-04.csv",
                          parse_dates=["data_coleta"])
        self.assertEqual(list(csv.columns), COLUNAS_COLETA)
        self.assertEqual(csv.loc[0, "preco"], 950)
        self.assertEqual(csv.loc[0, "data_coleta"],
                         pd.Timestamp(2026, 10, 4, 3))
        recuperada = coleta_para_carga([], dia=date(2026, 10, 4))
        self.assertEqual(len(recuperada), 1)
        self.assertEqual(executar_carga(recuperada)["ofertas"], 1)
        self.assertEqual(executar_carga(recuperada)["ofertas"], 0)


class TesteTrava(unittest.TestCase):
    """A trava impede duas coletas ao mesmo tempo, mas não pode ficar presa
    quando a execução anterior foi encerrada à força."""

    def setUp(self):
        import tempfile
        from pathlib import Path
        import main
        self.main = main
        self.pasta = tempfile.TemporaryDirectory()
        self.original = main.ARQUIVO_TRAVA
        main.ARQUIVO_TRAVA = Path(self.pasta.name) / ".pipeline.lock"

    def tearDown(self):
        self.main.ARQUIVO_TRAVA = self.original
        self.pasta.cleanup()

    @staticmethod
    def pid_encerrado():
        import subprocess
        import sys
        processo = subprocess.Popen([sys.executable, "-c", "pass"])
        processo.wait()
        return processo.pid

    def test_processo_ativo(self):
        import os
        self.assertTrue(self.main.processo_ativo(os.getpid()))
        self.assertFalse(self.main.processo_ativo(self.pid_encerrado()))

    def test_trava_de_execucao_encerrada_e_descartada(self):
        import os
        self.main.ARQUIVO_TRAVA.write_text(
            f"{self.pid_encerrado()};16/10/2026 18:12")
        self.assertIsNone(self.main.adquirir_trava())
        self.assertTrue(self.main.ARQUIVO_TRAVA.read_text().startswith(
            f"{os.getpid()};"))

    def test_trava_de_execucao_ativa_bloqueia(self):
        import os
        self.main.ARQUIVO_TRAVA.write_text(f"{os.getpid()};16/10/2026 18:12")
        self.assertEqual(self.main.adquirir_trava(),
                         (os.getpid(), "16/10/2026 18:12"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
